"""Reproduce the retained field transcription with bundled Python (read-only sources).

Headers in NSW examples take precedence over prose dictionary whitespace. PDF line
wraps in REC field names are joined with spaces. No eligibility rules are inferred.
"""
import json
import pathlib
import re

import openpyxl
import pdfplumber
from docx import Document

ROOT = pathlib.Path(__file__).parent


def clean(value):
    return re.sub(r"\s+", " ", value or "").strip()


def field(name, data_type, mandatory, validation, business="", reference=""):
    return dict(name=clean(name), dataType=clean(data_type), mandatory=clean(mandatory),
                validation=clean(validation), business=clean(business), reference=clean(reference))


manifest = json.loads((ROOT / "sources.json").read_text(encoding="utf-8"))


def source(filename):
    return next(item for item in manifest if item["file"] == filename)


schemas = {}
with pdfplumber.open(ROOT / "rec-sgu-swh-september-2025.pdf") as pdf:
    for key, pages, count in [("rec_sgu", range(7, 60), 146), ("rec_swh", range(60, 70), 65)]:
        fields = {}
        for page_number in pages:
            table = pdf.pages[page_number].extract_tables()[0]
            for row in table:
                # Two SWH pages contain extra merged number cells in the source.
                offset = 3 if key == "rec_swh" and page_number in (60, 69) else 1
                numbers = [cell for cell in row[:offset] if cell and cell.isdigit()]
                if not numbers:
                    continue
                number = int(numbers[0])
                f = field(*row[offset:offset + 5], reference=" ".join(c for c in row[offset + 5:-1] if c))
                f["sourceLocation"] = f"PDF page {page_number + 1}, field {number}"
                fields[number] = f
        assert sorted(fields) == list(range(1, count + 1)), (key, sorted(fields))
        schemas[key] = dict(fields=list(fields.values()), sources=[source("rec-sgu-swh-september-2025.pdf")], version="September 2025")

rows = [row for table in Document(ROOT / "rec-battery-fields.docx").tables for row in table.rows]
battery_fields = []
for row in rows:
    cells = [cell.text for cell in row.cells]
    if not re.fullmatch(r"\d+-[A-Z]+", cells[0]):
        continue
    f = field(*cells[1:7])
    f["sourceLocation"] = f"DOCX field {cells[0]}"
    battery_fields.append(f)
assert len(battery_fields) == 140
schemas["rec_battery"] = dict(fields=battery_fields, sources=[source("rec-battery-fields.docx"), source("rec-battery-fields.pdf"), source("rec-sgu-swh-september-2025.pdf")], version="Retrieved 24 September 2026")

specification = openpyxl.load_workbook(ROOT / "nsw-csv-v1.7.xlsx", data_only=True)
for key, sheet, template in [("nsw_esc", "ESS", "nsw-ess-example-v3.4.xlsx"), ("nsw_prc", "PDRS", "nsw-pdrs-example-v2.6.xlsx")]:
    dictionary = {clean(row[0]): (i, row) for i, row in enumerate(specification[sheet].values, 1) if row[0]}
    example = openpyxl.load_workbook(ROOT / template, data_only=True).worksheets[0]
    headers = [cell.value for cell in example[1]]
    fields = []
    for header in headers:
        row_number, row = dictionary[clean(header)]
        f = field(header, row[1], row[2], row[3])
        f["name"] = header  # Preserve exact official example header bytes.
        f["sourceLocation"] = f"Specification {sheet}!A{row_number}:E{row_number}; example row 1"
        fields.append(f)
    schemas[key] = dict(fields=fields, sources=[source("nsw-csv-v1.7.xlsx"), source(template)], version="TESSA CSV v1.7, 22 July 2026")

for key, schema in schemas.items():
    names = [f["name"] for f in schema["fields"]]
    assert len(set(names)) == len(names), key
    assert not any("\ufffd" in name for name in names), key

(ROOT / "schemas.json").write_text(json.dumps(schemas, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print({key: len(schema["fields"]) for key, schema in schemas.items()})
